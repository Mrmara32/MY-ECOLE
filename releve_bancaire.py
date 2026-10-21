"""
Lecture des relevés bancaires (CSV ou Excel .xlsx) pour le rapprochement bancaire.

Remplace pandas, qui n'était utilisé QUE pour cette lecture : une dépendance de
plusieurs dizaines de Mo en moins (installation et déploiements plus rapides,
exécutable Windows plus léger).

Au passage, les dates au format français JJ/MM/AAAA sont maintenant bien comprises
(pandas lisait 05/10/2026 comme le 10 mai, au format américain), et un fichier sans
ligne de titres n'en perd plus sa première opération.

Format attendu : 3 premières colonnes = Date, Description, Montant
(montant positif = crédit/entrée, négatif = débit/sortie).
"""
import csv
import io
from datetime import date, datetime

FORMATS_DATE = ('%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%Y-%m-%d', '%Y/%m/%d', '%d/%m/%y', '%d-%m-%y')


def lire_date(valeur):
    """Date au format AAAA-MM-JJ, ou None si la valeur n'est pas une date."""
    if isinstance(valeur, datetime):
        return valeur.strftime('%Y-%m-%d')
    if isinstance(valeur, date):
        return valeur.strftime('%Y-%m-%d')
    texte = str(valeur or '').strip()
    if not texte:
        return None
    texte = texte.split('T')[0].split(' ')[0]  # ignore une éventuelle heure
    for fmt in FORMATS_DATE:
        try:
            return datetime.strptime(texte, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue
    return None


def lire_montant(valeur):
    """Montant en nombre (négatif = débit), ou None si illisible.
    Accepte « 1 500 000 », « 1.500.000 », « 1,500,000 », « -250000 », « 250,50 »,
    « 1 500 000 GNF » et la notation comptable « (250 000) » pour un débit."""
    if isinstance(valeur, bool):
        return None
    if isinstance(valeur, (int, float)):
        return float(valeur)
    texte = str(valeur or '').strip()
    if not texte:
        return None
    negatif = texte.startswith('(') and texte.endswith(')')
    for a_retirer in (' ', ' ', ' ', 'GNF', 'gnf', 'FG', 'fg', '(', ')'):
        texte = texte.replace(a_retirer, '')
    if ',' in texte and '.' in texte:
        # Le dernier séparateur rencontré est celui des décimales, l'autre sépare les milliers.
        if texte.rfind(',') > texte.rfind('.'):
            texte = texte.replace('.', '').replace(',', '.')
        else:
            texte = texte.replace(',', '')
    elif texte.count(',') > 1:
        texte = texte.replace(',', '')
    elif texte.count('.') > 1:
        texte = texte.replace('.', '')
    else:
        texte = texte.replace(',', '.')
    try:
        montant = float(texte)
    except ValueError:
        return None
    return -abs(montant) if negatif else montant


def _lignes_csv(contenu):
    texte = None
    for encodage in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            texte = contenu.decode(encodage)
            break
        except UnicodeDecodeError:
            continue
    echantillon = texte[:4096]
    try:
        separateur = csv.Sniffer().sniff(echantillon, delimiters=';,\t|').delimiter
    except csv.Error:
        separateur = ';' if echantillon.count(';') >= echantillon.count(',') else ','
    return list(csv.reader(io.StringIO(texte), delimiter=separateur))


def _lignes_excel(contenu):
    from openpyxl import load_workbook
    classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    try:
        return [list(ligne) for ligne in classeur.worksheets[0].iter_rows(values_only=True)]
    finally:
        classeur.close()


def lignes_brutes(contenu, extension):
    """Toutes les lignes non vides du fichier (listes de cellules), sans interprétation.
    Renvoie (lignes, message_erreur)."""
    if extension == '.csv':
        brutes = _lignes_csv(contenu)
    elif extension == '.xlsx':
        brutes = _lignes_excel(contenu)
    else:
        return None, "Format non pris en charge : enregistrez le relevé au format .xlsx ou .csv"
    return [l for l in brutes if l and any(c is not None and str(c).strip() for c in l)], None


def lire_releve(contenu, extension):
    """Renvoie (lignes, message_erreur). Chaque ligne :
    {'date': 'AAAA-MM-JJ', 'description': str, 'montant': float positif, 'type': 'entree'|'sortie'}.
    Les lignes qui ne sont pas des opérations (titres, totaux, lignes vides) sont ignorées."""
    if extension == '.csv':
        brutes = _lignes_csv(contenu)
    elif extension == '.xlsx':
        brutes = _lignes_excel(contenu)
    else:
        return None, "Format non pris en charge : enregistrez le relevé au format .xlsx ou .csv"

    brutes = [l for l in brutes if l and any(c is not None and str(c).strip() for c in l)]
    if brutes and max(len(l) for l in brutes) < 3:
        return None, "Le fichier doit contenir au moins 3 colonnes : Date, Description, Montant"

    operations = []
    for brute in brutes:
        if len(brute) < 3:
            continue
        date_op, montant = lire_date(brute[0]), lire_montant(brute[2])
        if date_op is None or montant is None or montant == 0:
            continue
        operations.append({
            'date': date_op,
            'description': str(brute[1] if brute[1] is not None else '').strip(),
            'montant': abs(montant),
            'type': 'entree' if montant > 0 else 'sortie',
        })
    return operations, None
