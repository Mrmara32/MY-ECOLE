"""Relevé des dépenses par mois / catégorie (JSON et Excel)."""
import io


def _depense(client, entetes, date, montant, categorie):
    r = client.post('/api/transactions', json={'type': 'sortie', 'date_op': date, 'montant': montant,
                                                'categorie': categorie, 'description': f'Test {categorie}', 'journal': 'achats'},
                    headers=entetes)
    assert r.status_code == 201, r.get_json()


def test_releve_depenses_par_mois_et_categorie(client, entetes):
    _depense(client, entetes, '2031-01-10', 10000, 'Fournitures scolaires')
    _depense(client, entetes, '2031-01-25', 5000, 'Loyer')
    _depense(client, entetes, '2031-02-03', 7000, 'Fournitures scolaires')
    r = client.get('/api/releve/depenses?date_debut=2031-01-01&date_fin=2031-12-31', headers=entetes)
    assert r.status_code == 200, r.get_json()
    d = r.get_json()
    assert d['total'] == 22000 and d['nb'] == 3
    assert [m['mois'] for m in d['par_mois']] == ['2031-01', '2031-02']
    assert d['par_mois'][0]['total'] == 15000 and d['par_mois'][0]['categories']['Loyer'] == 5000
    assert {c['categorie']: c['total'] for c in d['par_categorie']} == {'Fournitures scolaires': 17000, 'Loyer': 5000}
    # La période est respectée
    d2 = client.get('/api/releve/depenses?date_debut=2031-02-01&date_fin=2031-02-28', headers=entetes).get_json()
    assert d2['total'] == 7000


def test_releve_depenses_excel(client, entetes):
    r = client.get('/api/releve/depenses/export?date_debut=2031-01-01&date_fin=2031-12-31', headers=entetes)
    assert r.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ['Par mois', 'Par catégorie', 'Détail']
